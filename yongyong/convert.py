#!/usr/bin/env python3
"""용용이가 만든 문서를 엑셀·워드·PDF 로 바꾼다 (대화창의 다운로드 버튼용).
읽기: .xlsx .csv .docx .md .txt  →  쓰기: .xlsx .docx .pdf (원본 옆에 같은 이름으로 저장)
  python3 convert.py 파일 xlsx|docx|pdf"""
import csv
import os
import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
SOURCES = {".xlsx", ".csv", ".docx", ".md", ".txt"}
TARGETS = ("xlsx", "docx", "pdf")


# ── 읽기: 문서를 [("title"|"p", 글자) / ("table", [[칸, ...], ...])] 블록 목록으로 ──
def read_blocks(path):
    ext = path.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        blocks = []
        for ws in wb.worksheets:
            rows = [["" if v is None else _cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(str(c).strip() for c in r)]
            if rows:
                width = max(len(r) for r in rows)
                while width and not any(len(r) >= width and str(r[width - 1]).strip() for r in rows):
                    width -= 1  # 오른쪽 빈 칸 정리
                if len(wb.worksheets) > 1:
                    blocks.append(("title", ws.title))
                blocks.append(("table", [list(r[:width]) + [""] * (width - len(r[:width])) for r in rows]))
        return blocks
    if ext == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [("table", [r for r in csv.reader(f) if any(c.strip() for c in r)])]
    if ext == ".docx":
        with zipfile.ZipFile(path) as z:
            body = ET.fromstring(z.read("word/document.xml")).find(W + "body")
        blocks = []
        for el in body:
            if el.tag == W + "p":
                text = "".join(t.text or "" for t in el.iter(W + "t")).strip()
                style = el.find(f"{W}pPr/{W}pStyle")
                is_title = style is not None and re.search(r"Heading|Title|제목", style.get(W + "val", ""), re.I)
                if text:
                    blocks.append(("title" if is_title else "p", text))
            elif el.tag == W + "tbl":
                rows = [["\n".join("".join(t.text or "" for t in p.iter(W + "t")) for p in tc.iter(W + "p")).strip()
                         for tc in tr.iter(W + "tc")] for tr in el.iter(W + "tr")]
                if rows:
                    blocks.append(("table", rows))
        return blocks
    # .md / .txt: 마크다운 표(| a | b |)는 표로, # 제목은 제목으로
    blocks, table = [], []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines() + [""]:
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            if not re.fullmatch(r"\|[\s:|-]+\|", s):
                table.append([c.strip() for c in s.strip("|").split("|")])
            continue
        if table:
            blocks.append(("table", table))
            table = []
        if s.startswith("#"):
            blocks.append(("title", s.lstrip("# ")))
        elif s:
            blocks.append(("p", re.sub(r"\*\*(.+?)\*\*", r"\1", s)))
    return blocks


def _cell(v):
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d") if not getattr(v, "hour", 0) and not getattr(v, "minute", 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


# ── 쓰기 ──
def to_xlsx(blocks, out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "문서"
    thin = Side(style="thin", color="BBBBBB")
    widths = {}
    r = 1
    for kind, data in blocks:
        if kind == "table":
            for i, row in enumerate(data):
                for c, v in enumerate(row, 1):
                    cell = ws.cell(r, c, v)
                    cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if i == 0:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="E8F5EE")
                    widths[c] = max(widths.get(c, 6), min(50, max((len(x) for x in str(v).split("\n")), default=0) * 2 + 2))
                r += 1
            r += 1
        else:
            ws.cell(r, 1, data).font = Font(bold=kind == "title", size=14 if kind == "title" else 11)
            r += 1
    for c, w in widths.items():
        ws.column_dimensions[ws.cell(1, c).column_letter].width = w
    wb.save(out)


def to_docx(blocks, out):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Apple SD Gothic Neo", Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Apple SD Gothic Neo")  # 한글 글꼴 (윈도에서는 맑은 고딕으로 대체됨)
    for kind, data in blocks:
        if kind == "title":
            doc.add_heading(data, level=1)
        elif kind == "p":
            doc.add_paragraph(data)
        else:
            cols = max(len(r) for r in data)
            t = doc.add_table(rows=0, cols=cols)
            t.style = "Table Grid"
            for i, row in enumerate(data):
                cells = t.add_row().cells
                for c in range(cols):
                    cells[c].text = str(row[c]) if c < len(row) else ""
                    if i == 0:
                        for run in cells[c].paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
    doc.save(out)


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for f in [os.environ.get("YONGYONG_PDF_FONT", ""), "/System/Library/Fonts/Supplemental/AppleGothic.ttf", "/Library/Fonts/NanumGothic.ttf",
              str(Path.home() / "Library/Fonts/NanumGothic.ttf"), "/usr/share/fonts/truetype/nanum/NanumGothic.ttf"]:
        if f and Path(f).exists():
            try:
                pdfmetrics.registerFont(TTFont("KR", f, subfontIndex=0))
                return "KR"
            except Exception:
                pass
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))  # 글꼴 파일이 없어도 되는 한글 글꼴
    return "HYGothic-Medium"


def to_pdf(blocks, out):
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    font = _pdf_font()
    wide = any(k == "table" and max(len(r) for r in d) > 5 for k, d in blocks)
    page = landscape(A4) if wide else A4
    body = ParagraphStyle("b", fontName=font, fontSize=9.5, leading=13, wordWrap="CJK")
    head = ParagraphStyle("h", parent=body, fontSize=15, leading=20, spaceAfter=6)
    para = lambda t, s=body: Paragraph(escape(str(t)).replace("\n", "<br/>"), s)  # noqa: E731
    story = []
    usable = page[0] - 30 * mm
    for kind, data in blocks:
        if kind == "table":
            cols = max(len(r) for r in data)
            rows = [[para(r[c] if c < len(r) else "") for c in range(cols)] for r in data]
            dw = lambda t: max((sum(2 if ord(ch) > 0x2E80 else 1.2 for ch in x) for x in str(t).split("\n")), default=0)  # noqa: E731  한글은 2칸
            lens = [max(6, max(min(dw(r[c]) if c < len(r) else 0, 60) for r in data) + 3) for c in range(cols)]
            widths = [usable * n / sum(lens) for n in lens]
            t = Table(rows, colWidths=widths, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BBBBBB")),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F5EE")),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story += [t, Spacer(1, 6 * mm)]
        else:
            story.append(para(data, head if kind == "title" else body))
    SimpleDocTemplate(str(out), pagesize=page, leftMargin=15 * mm, rightMargin=15 * mm,
                      topMargin=15 * mm, bottomMargin=15 * mm).build(story or [para(" ")])


def convert(path, fmt):
    """path 를 fmt(xlsx/docx/pdf) 로 바꾼 파일 경로. 이미 그 형식이면 원본, 원본이 더 새로우면 다시 만든다."""
    path = Path(path)
    if path.suffix.lower() == "." + fmt:
        return path
    if path.suffix.lower() not in SOURCES or fmt not in TARGETS:
        raise ValueError(f"{path.suffix} 파일은 {fmt} 로 바꿀 수 없어요.")
    out = path.with_suffix("." + fmt)
    if out.exists() and out.stat().st_mtime >= path.stat().st_mtime:
        return out
    {"xlsx": to_xlsx, "docx": to_docx, "pdf": to_pdf}[fmt](read_blocks(path), out)
    return out


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    print(convert(sys.argv[1], sys.argv[2]))
